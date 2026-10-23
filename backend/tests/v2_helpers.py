"""Shared test doubles for extraction v2: .docx builders and a fake
Anthropic client that records every request. Never touches the network.
"""

from __future__ import annotations

import io
import json
from types import SimpleNamespace
from typing import Any, Callable

from docx import Document
from docx.shared import Pt


def upes_course_docx() -> bytes:
    """The shapes that broke v1: bold "Step N." lines, a note after the
    steps, writer tags, a Linked_ line, a spec-wise fee table, multi-line cells."""
    d = Document()
    d.add_paragraph("Copy of UPES Online MBA program")
    d.add_heading("[about_heading] About UPES Online MBA", level=2)
    d.add_paragraph("UPES Online offers a two-year MBA. Established in 2003, it is NAAC A grade.")
    d.add_paragraph("Linked_university - UPES Online")
    d.add_heading("[admission_heading] UPES Online MBA Admission Process", level=2)
    d.add_paragraph("Follow these steps to apply:")
    for title, text in (("Step 1. Ensure your eligibility:", " Students are required to hold a bachelor's degree."),
                        ("Step 2. Register online:", " Create an account on the UPES Online portal.")):
        p = d.add_paragraph()
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(14)
        p.add_run(text)
    p = d.add_paragraph()
    p.add_run("Additional information:").bold = True
    p.add_run(" Keep scanned copies of your documents ready.")
    d.add_heading("[fee_heading] UPES Online MBA Specialization-wise Fees", level=2)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Specializations", "Fee"
    for name, fee in (("Marketing", "INR 25,000 /- Semester"), ("Sales and Marketing", "INR 40,000 /- Semester")):
        c = t.add_row().cells
        c[0].text, c[1].text = name, fee
    d.add_heading("Syllabus", level=2)
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Semester I", "Semester II"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Accounting\nEconomics", "Marketing\nFinance"
    # Subjects as a bulleted list inside a cell (UPES Marketing's syllabus table).
    t = d.add_table(rows=2, cols=1)
    t.rows[0].cells[0].text = "Semester III"
    cell = t.rows[1].cells[0]
    cell.paragraphs[0].text = "Consumer Behaviour"
    cell.paragraphs[0].style = d.styles["List Bullet"]
    cell.add_paragraph("Digital Brand Management", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


UPES_STEPS_HTML = ("<p>Follow these steps to apply:</p><ol>"
                   "<li><strong>Step 1. Ensure your eligibility:</strong> Students are required to hold a bachelor's degree.</li>"
                   "<li><strong>Step 2. Register online:</strong> Create an account on the UPES Online portal.</li></ol>")


def course_answer(**overrides: Any) -> dict[str, Any]:
    """A plausible model answer for upes_course_docx (extraction schema shape)."""
    from pipeline.v2.schema import extracted_fields

    values: dict[str, Any] = {}
    for f in extracted_fields("course"):
        values[f["key"]] = [] if f["type"] == "json" else ""
    values.update({
        "program_name": "UPES Online MBA",
        "university_name": "UPES Online",
        "naac_grade": "A",
        "about_heading": "About UPES Online MBA",
        "about_content": "<p>UPES Online offers a two-year MBA. Established in 2003, it is NAAC A grade.</p>",
        "admission_heading": "UPES Online MBA Admission Process",
        "admission_steps": UPES_STEPS_HTML,
        "admission_fee_note": "Additional information: Keep scanned copies of your documents ready.",
        "fee_heading": "",
        "emi_amount": "N/A",
        "specialization": [{"specialization_name": "Marketing", "specialization_details": ""},
                           {"specialization_name": "Sales and Marketing", "specialization_details": ""}],
        "syllabus_content": "<h4>Semester I</h4><ul><li>Accounting</li><li>Economics</li></ul>"
                            "<h4>Semester II</h4><ul><li>Marketing</li><li>Finance</li></ul>",
        "course_academic_years": [{"course_year_title": "Year I", "course_semesters": [
            {"course_semester_title": "Semester I", "course_semester_subjects": "<ul><li>Accounting</li><li>Economics</li></ul>"},
            {"course_semester_title": "Semester II", "course_semester_subjects": "<ul><li>Marketing</li><li>Finance</li></ul>"}]}],
    })
    values.update(overrides)
    provenance = [
        {"field": "program_name", "heading": "Copy of UPES Online MBA program", "confidence": "high"},
        {"field": "about_content", "heading": "[about_heading] About UPES Online MBA", "confidence": "high"},
        {"field": "admission_steps", "heading": "[admission_heading] UPES Online MBA Admission Process", "confidence": "high"},
        {"field": "admission_fee_note", "heading": "[admission_heading] UPES Online MBA Admission Process", "confidence": "medium"},
        {"field": "specialization", "heading": "UPES Online MBA Specialization-wise Fees", "confidence": "high"},
    ]
    return {"values": values, "provenance": provenance, "notes": ["No placement section."]}


GENERATED = {"seo_title": "UPES Online MBA: Fees, Admission and Syllabus | UPES Online",
             "meta_description": "UPES Online MBA: a two-year online MBA from UPES Online, NAAC A. See specializations, "
                                 "semester fees, admission steps and the full syllabus.",
             "hero_description": "A two-year online MBA from UPES Online with specializations in Marketing and Sales and Marketing."}


def message(data: Any, stop_reason: str = "end_turn", model: str = "claude-sonnet-5-5",
            usage: tuple[int, int] = (20000, 6000), extra_blocks: tuple = (), thinking: bool = True) -> SimpleNamespace:
    """thinking=False for v1's Haiku calls, whose replies start with the text block."""
    text = data if isinstance(data, str) else json.dumps(data)
    leading = [SimpleNamespace(type="thinking", thinking="")] if thinking else []
    return SimpleNamespace(
        content=[*leading, *extra_blocks, SimpleNamespace(type="text", text=text)],
        stop_reason=stop_reason, stop_details=None, model=model,
        usage=SimpleNamespace(input_tokens=usage[0], output_tokens=usage[1],
                              cache_read_input_tokens=0, cache_creation_input_tokens=0),
    )


class _Stream:
    def __init__(self, msg):
        self._msg = msg

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def get_final_message(self):
        return self._msg


class FakeAnthropic:
    """Records each stream() call; `respond(kwargs)` returns the message."""

    def __init__(self, respond: Callable[[dict[str, Any]], SimpleNamespace]):
        self.calls: list[dict[str, Any]] = []
        self._respond = respond
        self.beta = SimpleNamespace(messages=SimpleNamespace(stream=self._stream))

    def _stream(self, **kwargs):
        self.calls.append(kwargs)
        return _Stream(self._respond(kwargs))


def is_generation(kwargs: dict[str, Any]) -> bool:
    props = kwargs["output_config"]["format"]["schema"]["properties"]
    return "values" not in props


def standard_client(answer: dict[str, Any] | None = None, generated: dict[str, str] | None = None) -> FakeAnthropic:
    def respond(kwargs):
        if is_generation(kwargs):
            wanted = kwargs["output_config"]["format"]["schema"]["required"]
            return message({k: (generated or GENERATED).get(k, "") for k in wanted}, usage=(800, 300))
        return message(answer or course_answer())
    return FakeAnthropic(respond)
