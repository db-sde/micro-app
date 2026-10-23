"""Admission steps written as bold "Step N. Title:" lines, each followed by a
description paragraph (the UPES university document). The bold lines were
read as section headings, so the steps became separate untagged sections and
never reached admission_steps."""

from __future__ import annotations

import io

from docx import Document

from pipeline import wordpress_client as wc
from pipeline.docx_parser import parse_docx

STEPS = [
    ("Step 1. Ensure your eligibility:", "Check the eligibility criteria of the program."),
    ("Step 2. Fill out the application form:", "Click Apply Now and fill out the form."),
    ("Step 3. Make the payment:", "Pay the program fee to secure a seat."),
]


def upes_like_docx() -> bytes:
    doc = Document()
    doc.add_paragraph().add_run("[admission_heading,admission_steps] UPES Online Admission Process").bold = True
    doc.add_paragraph("Follow the below-listed admission process:")
    for title, body in STEPS:
        doc.add_paragraph().add_run(title).bold = True
        doc.add_paragraph(body)
    doc.add_paragraph("Additional information:")
    doc.add_paragraph("The application form fee is non-refundable.")
    doc.add_paragraph().add_run("[exam_content] Examination Process of UPES Online").bold = True
    doc.add_paragraph("Exams are held online through the LMS.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_bold_step_lines_stay_inside_the_admission_section():
    sections = parse_docx(upes_like_docx())["sections"]
    assert not any(h.startswith("step") for h in sections)
    admission = next(s for h, s in sections.items() if "admission" in h)["content"]
    for title, body in STEPS:
        assert title in admission and body in admission
    assert any("exam" in h for h in sections)


def test_each_step_keeps_its_description_and_trailing_info_becomes_the_note():
    html = "<p>Follow the process:</p>" + "".join(f"<p>{t}</p><p>{b}</p>" for t, b in STEPS) + \
           "<p>Additional information:</p><p>The application form fee is non-refundable.</p>"
    out = wc._polish_for_wordpress({"admission_steps": html}, "university")
    assert out["admission_steps"] == (
        "<p>Follow the process:</p><ol>"
        "<li><strong>Step 1. Ensure your eligibility:</strong> Check the eligibility criteria of the program.</li>"
        "<li><strong>Step 2. Fill out the application form:</strong> Click Apply Now and fill out the form.</li>"
        "<li><strong>Step 3. Make the payment:</strong> Pay the program fee to secure a seat.</li></ol>"
    )
    assert out["admission_fee_note"] == "The application form fee is non-refundable."


def test_step_titles_written_as_subheadings_with_an_info_list_after():
    # What extraction returned for the real UPES document.
    html = ("<p>Follow the process:</p>"
            "<h4>Step 1. Ensure your eligibility:</h4><p>Check the criteria.</p>"
            "<h4>Step 2. Make the payment:</h4><p>Pay the fee.</p>"
            "<h4>Additional information:</h4><ul><li>The form fee is non-refundable.</li>"
            "<li>Note the transaction number.</li></ul>")
    out = wc._polish_for_wordpress({"admission_steps": html,
                                    "admission_fee_note": "Step 2. Make the payment: Pay the fee."}, "university")
    assert out["admission_steps"] == (
        "<p>Follow the process:</p><ol>"
        "<li><strong>Step 1. Ensure your eligibility:</strong> Check the criteria.</li>"
        "<li><strong>Step 2. Make the payment:</strong> Pay the fee.</li></ol>"
    )
    assert out["admission_fee_note"] == "The form fee is non-refundable. Note the transaction number."
