"""The .docx as structured Markdown, for v2's single extraction call.

Docling's Word backend only (no OCR, layout or ML models). Heading levels,
ordered/unordered/nested lists, bold text and full tables are kept.
Export settings, all checked against our documents' shapes (see
EXTRACTION_V2_PLAN.md §6a):

* escape_underscores=False keeps writer tags ([about_heading]) and
  "Linked_university - …" lines verbatim;
* escape_html=False keeps "Larsen & Toubro" instead of "&amp;";
* line breaks inside table cells become <br>, so multi-line cells (one
  subject per line) don't run together.

Only Word heading styles become Markdown headings. A bold
"Step 1. Ensure your eligibility:" line stays a bold line inside its
section; v1's heading detector split those into sections, which is how
the UPES admission steps were lost.

If Docling fails or isn't installed, the document goes through the
existing docx_parser instead (parser="fallback").
"""

from __future__ import annotations

import io
import logging
import threading
from dataclasses import dataclass, field
from typing import Any

logger = logging.getLogger("degreebaba.v2.markdown")

_converter = None
_converter_lock = threading.Lock()


@dataclass
class MarkdownDoc:
    markdown: str
    tables: list[list[list[str]]] = field(default_factory=list)  # every table: rows of cell text
    title: str = ""            # first line of the document (before the first heading)
    parser: str = "docling"    # "docling" or "fallback"
    error: str | None = None   # why the fallback was used


def to_markdown(file_bytes: bytes, filename: str = "document.docx") -> MarkdownDoc:
    try:
        return _docling(file_bytes, filename)
    except Exception as exc:
        logger.warning("DOCLING_FAILED: %s (%s); using docx_parser", filename, exc)
        doc = _fallback(file_bytes)
        doc.error = f"{type(exc).__name__}: {exc}"
        return doc


def _get_converter():
    global _converter
    with _converter_lock:
        if _converter is None:
            from docling.datamodel.base_models import InputFormat
            from docling.document_converter import DocumentConverter
            _converter = DocumentConverter(allowed_formats=[InputFormat.DOCX])
        return _converter


def _docling(file_bytes: bytes, filename: str) -> MarkdownDoc:
    from docling.datamodel.base_models import DocumentStream

    name = filename if filename.lower().endswith(".docx") else f"{filename}.docx"
    converter = _get_converter()
    with _converter_lock:  # one conversion at a time; each takes milliseconds
        result = converter.convert(DocumentStream(name=name, stream=io.BytesIO(file_bytes)))
    doc = result.document

    from docling_core.types.doc import TableCell

    tables: list[list[list[str]]] = []
    cell_groups = []
    for table in doc.tables:
        tables.append([[(cell.text or "").strip() for cell in row] for row in table.data.grid])
        cells = table.data.table_cells
        for i, cell in enumerate(cells):
            # A cell holding a list is a "rich" cell pointing at a list group, which
            # the export renders inline ("- A - B - C"), so one subject per line is
            # lost. Use a plain cell with <br> between lines, and drop the group.
            ref = getattr(cell, "ref", None)
            if ref is not None:
                cell_groups.append(ref.resolve(doc))
            plain = TableCell.model_validate(cell.model_dump(exclude={"ref"}))
            plain.text = (cell.text or "").strip().replace("\n", "<br>")
            cells[i] = plain
    if cell_groups:
        doc.delete_items(node_items=cell_groups)

    markdown = doc.export_to_markdown(escape_underscores=False, escape_html=False)
    title = next((line.strip() for line in markdown.splitlines() if line.strip()), "")
    if title.startswith("#"):
        title = title.lstrip("#").strip()
    return MarkdownDoc(markdown=markdown.strip() + "\n", tables=tables, title=title, parser="docling")


# ── fallback: the existing parser ──────────────────────────────────────

def _fallback(file_bytes: bytes) -> MarkdownDoc:
    from docx import Document

    from pipeline.docx_parser import parse_docx

    parsed = parse_docx(file_bytes)
    meta = parsed.get("__meta__", {})
    lines: list[str] = []
    title = meta.get("page_title") or ""
    if title:
        lines += [title, ""]
    for section in parsed.get("sections", {}).values():
        heading = section.get("original_heading") or section.get("stripped_key") or ""
        if heading and heading != title:
            lines += [f"## {heading}", ""]
        lines += [_render(section.get("content")), ""]

    tables = []
    for table in Document(io.BytesIO(file_bytes)).tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        rows = [r for r in rows if any(r)]
        if rows:
            tables.append(rows)
    markdown = "\n".join(lines).strip() + "\n"
    return MarkdownDoc(markdown=markdown, tables=tables, title=title, parser="fallback")


def _table(rows: list[dict[str, Any]]) -> str:
    headers = list(rows[0])
    out = ["| " + " | ".join(h.replace("_", " ") for h in headers) + " |",
           "|" + "---|" * len(headers)]
    for row in rows:
        out.append("| " + " | ".join(str(row.get(h, "")).replace("\n", "<br>") for h in headers) + " |")
    return "\n".join(out)


def _render(content: Any) -> str:
    if content is None:
        return ""
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        if content and all(isinstance(x, dict) and "question" in x for x in content):
            return "\n\n".join(f"**{x.get('question', '')}**\n\n{x.get('answer', '')}" for x in content)
        if content and all(isinstance(x, dict) for x in content):
            return _table(content)
        return "\n".join(f"- {x}" for x in content)
    if isinstance(content, dict):
        parts = []
        if content.get("intro_text"):
            parts.append(str(content["intro_text"]))
        for table in content.get("tables") or []:
            parts.append(_render((table or {}).get("data")))
        if parts:
            return "\n\n".join(p for p in parts if p)
        # semester / horizontal tables: key -> value(s)
        lines = []
        for key, value in content.items():
            if isinstance(value, dict):
                lines.append(f"### {key}")
                for sub, items in value.items():
                    items = items if isinstance(items, list) else [items]
                    lines.append(f"#### {sub}\n" + "\n".join(f"- {i}" for i in items))
            elif isinstance(value, list):
                lines.append(f"**{key}**\n" + "\n".join(f"- {i}" for i in value))
            else:
                lines.append(f"**{key}:** {value}")
        return "\n\n".join(lines)
    return str(content)
